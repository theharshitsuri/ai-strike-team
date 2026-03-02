"""
core/ingestion.py — File and email ingestion helpers.
Handles PDF, DOCX, Excel, CSV, and raw text extraction.

Usage:
    from core.ingestion import ingest_file, fetch_emails
    text = ingest_file("path/to/invoice.pdf")
"""

import email
import imaplib
from pathlib import Path
from typing import Generator

from core.config import settings
from core.logger import get_logger

log = get_logger(__name__)


# ─── File Ingestion ─────────────────────────────────────────────────────────

def ingest_file(path: str | Path) -> str:
    """
    Ingest a file and return its text content.
    Supports: .pdf, .docx, .xlsx, .csv, .txt
    """
    path = Path(path)
    suffix = path.suffix.lower()
    log.info("ingesting_file", path=str(path), type=suffix)

    if suffix == ".pdf":
        return _ingest_pdf(path)
    elif suffix == ".docx":
        return _ingest_docx(path)
    elif suffix in (".xlsx", ".xls"):
        return _ingest_excel(path)
    elif suffix == ".csv":
        return _ingest_csv(path)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _ingest_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _ingest_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _ingest_excel(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join(str(c) if c is not None else "" for c in row))
    return "\n".join(lines)


def _ingest_csv(path: Path) -> str:
    return path.read_text(encoding="utf-8")


# ─── Email Ingestion ─────────────────────────────────────────────────────────

def fetch_emails(
    folder: str = "INBOX",
    unseen_only: bool = True,
    limit: int = 20,
) -> Generator[dict, None, None]:
    """
    Fetch emails from IMAP and yield dicts with subject, sender, body, date.
    Set IMAP credentials in .env.
    """
    log.info("fetching_emails", folder=folder, unseen_only=unseen_only, limit=limit)
    mail = imaplib.IMAP4_SSL(settings.email_imap_host, settings.email_imap_port)
    mail.login(settings.email_address, settings.email_password)
    mail.select(folder)

    search_criterion = "(UNSEEN)" if unseen_only else "ALL"
    _, message_ids = mail.search(None, search_criterion)
    ids = message_ids[0].split()[-limit:]  # most recent N

    for uid in ids:
        _, msg_data = mail.fetch(uid, "(RFC822)")
        raw = msg_data[0][1]
        msg = email.message_from_bytes(raw)

        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    body = part.get_payload(decode=True).decode("utf-8", errors="replace")
                    break
        else:
            body = msg.get_payload(decode=True).decode("utf-8", errors="replace")

        yield {
            "uid": uid.decode(),
            "subject": msg.get("Subject", ""),
            "sender": msg.get("From", ""),
            "date": msg.get("Date", ""),
            "body": body,
        }

    mail.logout()
